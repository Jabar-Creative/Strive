"""Test ekstraksi PDF/DOCX + penyusunan dari profil — AI-02.

AC AI-02: "CV lama dua kolom tetap terbaca. Dokumen tanpa teks (hasil scan
gambar) ditolak dengan pesan yang jelas."

Semua berkas uji DIBUAT di sini (reportlab menulis PDF, python-docx menulis
DOCX, pypdf mengenkripsi) — bukan hasil rekaman layar — supaya test bisa
diperbaiki bersama perilaku yang diujinya, dan fixture-nya jujur: PDF dua
kolom benar-benar digambar dua kolom, bukan teks yang diklaim dua kolom.

Kode galat yang dites adalah daftar TERTUTUP §10.2 (isu #197):
FILE_TOO_LARGE · UNSUPPORTED_FILE_TYPE · DOCUMENT_TOO_LONG · SOURCE_TOO_SHORT.
"""

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from app.core.errors import GalatLayanan
from app.pipelines.extract import ProfilCv, ekstrak_dokumen, susun_dari_profil

# ── Fixture berkas ────────────────────────────────────────────────────────


def buat_pdf_satu_kolom() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    baris_cv = [
        "Budi Santoso",
        "budi.santoso@yopmail.com",
        "Pengalaman kerja:",
        "Software Engineer PT Maju 2021-06 sampai 2024-02",
        "Membangun API pembayaran untuk 2 juta pengguna aktif",
        "Menulis 40 unit test dan menurunkan bug produksi 30 persen",
        "Keterampilan: Python TypeScript PostgreSQL Redis",
    ]
    y = 780
    for baris in baris_cv:
        c.drawString(72, y, baris)
        y -= 24
    c.save()
    return buf.getvalue()


def buat_pdf_dua_kolom() -> bytes:
    """AC utama: teks kiri harus terbaca UTUH sebelum teks kanan."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    c.setFont("Helvetica", 10)
    kiri = [
        "Kiri pengalaman kerja satu",
        "Kiri pengalaman kerja dua",
        "Kiri pengalaman kerja tiga",
        "Kiri pengalaman kerja empat",
        "Kiri pengalaman kerja lima",
        "Kiri pengalaman kerja enam",
    ]
    kanan = [
        "Kanan pendidikan gelar sarjana",
        "Kanan pendidikan magister teknik",
        "Kanan sertifikasi AWS tersedia",
        "Kanan sertifikasi GCP tersedia",
        "Kanan kontak email yopmail",
        "Kanan kontak telepon aktif",
    ]
    y = 780
    for i in range(6):
        c.drawString(40, y, kiri[i])
        c.drawString(330, y, kanan[i])
        y -= 24
    c.save()
    return buf.getvalue()


def buat_pdf_dua_kolom_tajuk() -> bytes:
    """CV klasik: tajuk lebar penuh + dua kolom dengan panjang BERBEDA.

    Regresi verifikasi nyata: versi pertama salah memindahkan baris kiri
    ke tajuk saat sisinya cuma satu kata ("KETERAMPILAN"), dan menggabung
    judul seksi "PENGALAMAN | PENDIDIKAN" — urutan baca rusak diam-diam
    padahal fixture sintetis (4+4 kata per baris) hijau.
    """
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, 800, "Dewi Lestari")
    c.setFont("Helvetica", 10)
    c.drawString(40, 784, "dewi.lestari@yopmail.com · Bandung")
    c.setFont("Helvetica", 9)
    kiri = [
        "PENGALAMAN",
        "Product Designer, PT Kreasi (2022-04 - kini)",
        "Memimpin desain ulang checkout; konversi naik 18 persen.",
        "Membangun design system dipakai 4 tim produk.",
        "UI Designer, CV Media (2020-01 - 2022-03)",
        "Merancang 30+ layar aplikasi berita mobile.",
    ]
    kanan = [
        "PENDIDIKAN",
        "S1 Desain Komunikasi Visual, ITB (2019)",
        "KETERAMPILAN",
        "Figma, user research, prototyping,",
        "usability testing, design tokens,",
        "kolaborasi lintas tim yang rapat.",
    ]
    y = 750
    for i in range(6):
        c.drawString(40, y, kiri[i])
        c.drawString(320, y, kanan[i])
        y -= 18
    c.save()
    return buf.getvalue()


def buat_pdf_tanpa_teks() -> bytes:
    """Halaman berisi gambar bingkai saja — meniru hasil scan gambar."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    c.rect(100, 500, 395, 300)
    c.showPage()
    c.save()
    return buf.getvalue()


def buat_pdf_terenkripsi() -> bytes:
    masuk = io.BytesIO(buat_pdf_satu_kolom())
    pembaca = PdfReader(masuk)
    penulis = PdfWriter()
    for halaman in pembaca.pages:
        penulis.add_page(halaman)
    penulis.encrypt("rahasia-tes")
    keluar = io.BytesIO()
    penulis.write(keluar)
    return keluar.getvalue()


def buat_docx_cv() -> bytes:
    dok = Document()
    dok.add_paragraph("Sinta Wijaya")
    dok.add_paragraph("sinta.wijaya@yopmail.com")
    dok.add_paragraph("Data analyst tiga tahun di sektor fintech dengan fokus")
    dok.add_paragraph("pada eksperimen pertumbuhan dan dashboard eksekutif.")
    tabel = dok.add_table(rows=2, cols=2)
    tabel.cell(0, 0).text = "Python"
    tabel.cell(0, 1).text = "Lanjutan"
    tabel.cell(1, 0).text = "SQL"
    tabel.cell(1, 1).text = "Mahir"
    buf = io.BytesIO()
    dok.save(buf)
    return buf.getvalue()


def buat_pdf_banyak_halaman(jumlah: int) -> bytes:
    """Halaman kosong bertumpuk — objek halaman kecil, kata nol.

    Meniru temuan audit: PDF bisa memuat ratusan ribu objek halaman yang
    masing-masing dijalankan layout analysis; cap harus menyala SEBELUM
    iterasi, bukan setelah seluruh pekerjaan mahal selesai.
    """
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    for _ in range(jumlah):
        c.showPage()
    c.save()
    return buf.getvalue()


def buat_pdf_padat_kata(jumlah_kata: int) -> bytes:
    """Kata padat lintas halaman — anggaran kata harus menyala di tengah jalan."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    c.setFont("Helvetica", 8)
    per_baris = 12
    y = 800
    for _ in range(0, jumlah_kata, per_baris):
        c.drawString(30, y, ("kata " * per_baris).strip())
        y -= 11
        if y < 40:
            c.showPage()
            c.setFont("Helvetica", 8)
            y = 800
    c.save()
    return buf.getvalue()


def buat_zip_biasa() -> bytes:
    """Magic PK benar (zip), tapi bukan DOCX — jebakan K-01 berulang."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("catatan.txt", "ini zip biasa, bukan docx")
    return buf.getvalue()


def buat_docx_panjang(jumlah_kata: int) -> bytes:
    dok = Document()
    per_paragraf = 50
    for _ in range(jumlah_kata // per_paragraf):
        dok.add_paragraph("kata " * per_paragraf)
    buf = io.BytesIO()
    dok.save(buf)
    return buf.getvalue()


def buat_docx_bom() -> bytes:
    """ZIP sah bentuknya (terkompresi kecil), tapi memuai saat dibuka."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as z:
        # 2 × 80 MB nol terurai; terkompresi puluhan KB saja. Total terurai
        # 160 MB melampaui ambang dekompresi 100 MB sebelum isi dibaca.
        z.writestr("entri-besar-1.bin", b"\x00" * (80 * 1024 * 1024))
        z.writestr("entri-besar-2.bin", b"\x00" * (80 * 1024 * 1024))
    return buf.getvalue()


# ── Ekstraksi: jalur bahagia ──────────────────────────────────────────────


def test_pdf_satu_kolom_terbaca() -> None:
    hasil = ekstrak_dokumen(buat_pdf_satu_kolom())
    assert hasil.sumber == "pdf"
    assert "Budi Santoso" in hasil.teks
    assert "PT Maju" in hasil.teks
    # Satu kolom: tidak ada temuan tata letak.
    assert not any("dua kolom" in t.pesan for t in hasil.temuan)


def test_pdf_dua_kolom_terbaca_utuh_dan_ditandai() -> None:
    hasil = ekstrak_dokumen(buat_pdf_dua_kolom())
    teks = hasil.teks
    # Kolom kiri UTUH sebelum kolom kanan dimulai — bukan tercampur
    # per baris visual seperti extract_text polos.
    assert teks.index("Kiri pengalaman kerja enam") < teks.index("Kanan pendidikan gelar")
    # Tidak ada baris yang menggabungkan kiri dan kanan.
    for baris in teks.splitlines():
        if "Kiri" in baris:
            assert "Kanan" not in baris, f"baris tercampur: {baris!r}"
    # PRD E8 kasus tepi: "ditandai di findings".
    assert any("dua kolom" in t.pesan.lower() for t in hasil.temuan)


def test_pdf_dua_kolom_dengan_tajuk_urutan_baca_utuh() -> None:
    hasil = ekstrak_dokumen(buat_pdf_dua_kolom_tajuk())
    teks = hasil.teks
    # Tajuk halaman tetap di atas.
    assert teks.startswith("Dewi Lestari")
    # Kolom kiri UTUH dan berurutan, semua sebelum isi kanan pertama.
    for kiri_lebih_dulu in [
        "Product Designer, PT Kreasi",
        "Memimpin desain ulang checkout",
        "Membangun design system",
        "UI Designer, CV Media",
        "Merancang 30+ layar",
    ]:
        assert teks.index(kiri_lebih_dulu) < teks.index("S1 Desain Komunikasi Visual"), (
            f"baris kiri {kiri_lebih_dulu!r} jatuh setelah kolom kanan"
        )
    # Baris kiri TIDAK tercampur dengan kata tunggal kanan di tajuk.
    assert not any(
        "Memimpin" in baris and "KETERAMPILAN" in baris for baris in teks.splitlines()
    )
    # Tetap ditandai dua kolom.
    assert any("dua kolom" in t.pesan.lower() for t in hasil.temuan)


def test_docx_paragraf_dan_tabel_terbaca() -> None:
    hasil = ekstrak_dokumen(buat_docx_cv())
    assert hasil.sumber == "docx"
    assert "Sinta Wijaya" in hasil.teks
    assert "Data analyst" in hasil.teks
    # Isi tabel ikut terbaca — CV lama sering memakai tabel untuk skill.
    assert "Python" in hasil.teks
    assert "Mahir" in hasil.teks


# ── Ekstraksi: penolakan dengan pesan jelas ───────────────────────────────


def test_pdf_tanpa_teks_ditolak_pesan_scan() -> None:
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(buat_pdf_tanpa_teks())
    assert info.value.code == "SOURCE_TOO_SHORT"
    # Pesan harus MENYEBUT kemungkinan scan gambar — "pesan jelas" di AC.
    assert "scan" in info.value.message.lower()


def test_pdf_terenkripsi_ditolak_pesan_password() -> None:
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(buat_pdf_terenkripsi())
    assert info.value.code == "UNSUPPORTED_FILE_TYPE"
    assert "password" in info.value.message.lower() or "enkripsi" in info.value.message.lower()


def test_bukan_pdf_docx_ditolak() -> None:
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(b"<html><body>ini bukan dokumen</body></html>")
    assert info.value.code == "UNSUPPORTED_FILE_TYPE"


def test_zip_biasa_bukan_docx_ditolak() -> None:
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(buat_zip_biasa())
    assert info.value.code == "UNSUPPORTED_FILE_TYPE"


def test_dokumen_lebih_25mb_ditolak() -> None:
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(b"%PDF-1.4\n" + b"\x00" * (25 * 1024 * 1024 + 1))
    assert info.value.code == "FILE_TOO_LARGE"
    assert info.value.status == 413


def test_dokumen_diatas_15000_kata_ditolak() -> None:
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(buat_docx_panjang(15_050))
    assert info.value.code == "DOCUMENT_TOO_LONG"
    assert info.value.status == 422


def test_pdf_diatas_200_halaman_ditolak_sebelum_iterasi() -> None:
    # Temuan audit: cap halaman harus menyala SEBELUM layout analysis,
    # bukan setelah seluruh dokumen diproses.
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(buat_pdf_banyak_halaman(201))
    assert info.value.code == "DOCUMENT_TOO_LONG"
    assert "halaman" in info.value.message


def test_pdf_anggaran_kata_menyala_di_tengah_dokumen() -> None:
    # Padat kata lintas halaman (di bawah cap halaman): penolakan tetap
    # terjadi dari akumulasi DALAM loop, tidak menunggu dokumen habis.
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(buat_pdf_padat_kata(15_050))
    assert info.value.code == "DOCUMENT_TOO_LONG"


def test_docx_bom_dekompresi_ditolak() -> None:
    """Zip sah ≤25 MB terkompresi, tapi memuai melampaui batas saat dibuka."""
    with pytest.raises(GalatLayanan) as info:
        ekstrak_dokumen(buat_docx_bom())
    assert info.value.code == "FILE_TOO_LARGE"


# ── Penyusunan dari profil ────────────────────────────────────────────────


def profil_lengkap() -> ProfilCv:
    return ProfilCv(
        nama="Rina Melati",
        headline="Backend engineer yang suka sistem terdistribusi",
        kontak={"email": "rina.melati@yopmail.com"},
        pengalaman=[
            {
                "posisi": "Backend Engineer",
                "perusahaan": "PT Nusantara Digital",
                "mulai": "2022-03",
                "selesai": "2025-08",
                "deskripsi": "Merancang layanan pembayaran dengan throughput 500 rps.",
            }
        ],
        pendidikan=[{"gelar": "S1 Informatika", "institusi": "Universitas Gadjah Mada"}],
        keterampilan=["Go", "PostgreSQL", "Kafka"],
    )


def test_susun_dari_profil_memuat_semua_bagian() -> None:
    teks = susun_dari_profil(profil_lengkap())
    for bagian in ("Rina Melati", "PT Nusantara Digital", "2022-03", "S1 Informatika", "Kafka"):
        assert bagian in teks


def test_susun_dari_profil_deterministik() -> None:
    assert susun_dari_profil(profil_lengkap()) == susun_dari_profil(profil_lengkap())


def test_susun_dari_profil_terlalu_pendek_ditolak() -> None:
    with pytest.raises(GalatLayanan) as info:
        susun_dari_profil(ProfilCv(nama="Rina Melati"))
    assert info.value.code == "SOURCE_TOO_SHORT"
